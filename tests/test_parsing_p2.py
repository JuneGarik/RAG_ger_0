"""Этап П2: тепловые карты и подписи диаграмм кодом, координаты групп, проверка ответа модели."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest

pptx_lib = pytest.importorskip("pptx")

from bzpo.parsing import heatmap, vlm  # noqa: E402
from bzpo.parsing.pptx import (  # noqa: E402
    HEATMAP_NOTE,
    LEGEND_CHART_NOTE,
    UNNAMED_CHART_NOTE,
    SlideData,
    _flatten,
    extract,
    to_markdown,
)

# --------------------------------------------------------------- тепловая карта

LEGEND = [heatmap.LegendItem("N/A", heatmap.THEME), heatmap.LegendItem("≥ 2.4", "E5FAFB"),
          heatmap.LegendItem("2-2.1", "B4EEF1"), heatmap.LegendItem("≤ 1.7", "02C0CA")]


def test_heatmap_maps_shades_by_brightness_order() -> None:
    rows = [
        ["", "", ""],  # шапка, залитая цветом темы, — не «N/A»
        ["Видение", "■ #4ED1D8", "■ #54CFD6"],  # тёмные оттенки ≈ ≤ 1.7
        ["Данные", "■ #B3EDF0", "■ THEME"],  # средний; THEME в строке значений -> N/A
        ["ИТ", "■ #E4F9FA", "■ #E5FAFB"],  # светлые ≈ ≥ 2.4
    ]
    rows[0][1] = "■ THEME"
    mapped = heatmap.map_cells(rows, LEGEND)
    assert mapped == [["", "", ""], ["Видение", "≤ 1.7", "≤ 1.7"],
                      ["Данные", "2-2.1", "N/A"], ["ИТ", "≥ 2.4", "≥ 2.4"]]


def test_heatmap_refuses_when_shades_do_not_match_legend() -> None:
    rows = [["Видение", "■ #4ED1D8", "■ #B3EDF0", "■ #E4F9FA", "■ #7F7F7F"]]
    assert heatmap.map_cells(rows, LEGEND) is None  # 4 оттенка на 3 цветных позиции


# --------------------------------------------------------------- диаграмма с легендой


def _box(slide: Any, x: float, y: float, w: float, h: float, text: str = "",
         rgb: str | None = None) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w),
                                   Inches(h))
    if rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(rgb)
    else:
        shape.fill.background()
    shape.line.fill.background()
    if text:
        shape.text_frame.text = text
    return shape


def build_think_cell_like(path: Path) -> Path:
    """Диаграмма без имён рядов и категорий в XML, подписи — надписи на слайде."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Выручка снизилась"
    data = CategoryChartData()
    data.categories = ["", ""]
    data.add_series("", (146764.95, 128649.79))
    data.add_series("", (45765.5, 38023.17))
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5),
                                         Inches(2.5), Inches(6), Inches(3.5), data)
    for series, color in zip(chart_shape.chart.plots[0].series, ("02CAD4", "9AF7FF"),
                             strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    # Как в think-cell: ни категорий, ни имён рядов в XML диаграммы нет.
    space = chart_shape.chart._chartSpace
    for element in space.xpath(".//c:cat") + space.xpath(".//c:ser/c:tx"):
        element.getparent().remove(element)
    # Легенда над диаграммой: маркер того же цвета + подпись справа.
    _box(slide, 0.6, 1.6, 0.2, 0.15, rgb="9AF7FF")
    _box(slide, 0.9, 1.57, 1.5, 0.2, "EBITDA")
    _box(slide, 0.6, 1.9, 0.2, 0.15, rgb="02CAD4")
    _box(slide, 0.9, 1.87, 1.5, 0.2, "Выручка")
    # Подписи категорий — в нижней полосе диаграммы.
    _box(slide, 1.5, 5.7, 0.6, 0.2, "2023")
    _box(slide, 4.0, 5.7, 0.6, 0.2, "2024")
    target = path / "chart.pptx"
    prs.save(str(target))
    return target


def test_chart_labels_from_legend_colors(tmp_path: Path) -> None:
    md = to_markdown(extract(build_think_cell_like(tmp_path)))
    assert LEGEND_CHART_NOTE in md and UNNAMED_CHART_NOTE not in md
    # Порядок рядов — порядок данных, а не порядок строк легенды.
    assert "| Категория | Выручка | EBITDA |" in md
    assert "| 2023 | 146764,95 | 45765,5 |" in md
    assert "| 2024 | 128649,79 | 38023,17 |" in md


def test_group_children_get_slide_coordinates(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Emu(100), Emu(200), Emu(1000), Emu(500)).text_frame.text = "А"
    xfrm = group._element.grpSpPr.get_or_add_xfrm()
    xfrm.get_or_add_off().x, xfrm.get_or_add_off().y = 5000, 7000
    xfrm.get_or_add_ext().cx, xfrm.get_or_add_ext().cy = 2000, 1000
    xfrm.get_or_add_chOff().x, xfrm.get_or_add_chOff().y = 100, 200
    xfrm.get_or_add_chExt().cx, xfrm.get_or_add_chExt().cy = 1000, 500
    path = tmp_path / "group.pptx"
    prs.save(str(path))
    [shape] = _flatten(Presentation(str(path)).slides[0].shapes)
    # Ребёнок в (100, 200) системы группы = левый верхний угол группы на слайде, масштаб ×2.
    assert (shape.left, shape.top, shape.width, shape.height) == (5000, 7000, 2000, 1000)


# --------------------------------------------------------------- проверка ответа модели


class FakeVision:
    def __init__(self, answer: dict[str, Any]) -> None:
        self.answer = answer
        self.calls = 0

    def generate(self, prompt: str, **kwargs: Any) -> Any:
        self.calls += 1
        assert kwargs["images"] and kwargs["json_format"]

        class Result:
            text = json.dumps(self.answer, ensure_ascii=False)

        return Result()


def slide() -> SlideData:
    data = SlideData(16, False, "Выручка снизилась на 12% до 129 млрд", "01. Финансы")
    data.blocks = ["Финансовые показатели компании", "*[Изображение: схема — ...]*"]
    data.images = ["схема"]
    return data


def test_vision_answer_verified_and_cached(tmp_path: Path) -> None:
    image = tmp_path / "page.png"
    image.write_bytes(b"png")
    client = FakeVision({
        "key_message": "Выручка снизилась на 12% до 129 млрд, маржа упала на 7 п.п. до 31%",
        "visual": "Столбцы по годам",
        "image_text": "Финансовые показатели компании",  # повтор извлечённого текста
    })
    result = vlm.analyze(slide(), image, client, tmp_path / "cache", "model")
    assert result.key_message.startswith("Выручка снизилась")
    assert any("31" in issue for issue in result.issues)  # числа 31 нет на слайде
    assert result.image_text == ""  # повтор отброшен
    again = vlm.analyze(slide(), image, client, tmp_path / "cache", "model")
    assert again.cached and client.calls == 1


def test_vision_noise_and_apply(tmp_path: Path) -> None:
    image = tmp_path / "page.png"
    image.write_bytes(b"png2")
    client = FakeVision({"key_message": "Вывод", "visual": "Схема",
                         "image_text": "COMPANY"})  # слово с фотографии — шум
    data = slide()
    result = vlm.analyze(data, image, client, tmp_path / "cache", "model")
    assert result.image_text == ""
    client.answer["image_text"] = "Платформа данных связывает продажи, стройку и финансы"
    image.write_bytes(b"png3")
    result = vlm.analyze(data, image, client, tmp_path / "cache", "model")
    vlm.apply(data, result)
    data.vlm_done = True
    md = to_markdown_one(data)
    assert "**Ключевой посыл:** Вывод" in md and "**Визуал:** Схема" in md
    assert "**Текст с изображения (распознан моделью, проверить):** Платформа" in md
    assert "изображение разобрано моделью" in md


def to_markdown_one(data: SlideData) -> str:
    from bzpo.parsing.pptx import DeckData

    return to_markdown(DeckData(Path("x.pptx"), "Т", [data]))


def test_heatmap_note_constant_present() -> None:
    assert "по цвету заливки и легенде" in HEATMAP_NOTE
