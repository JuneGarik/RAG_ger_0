"""Черновик строки реестра: разбор ответа LLM, проверки по слайдам, роли, запись в Excel."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest

from bzpo.ingest.slides import chunk_slides
from bzpo.registry.draft import (
    DraftContext,
    apply_team_table,
    parse,
    select_context,
    tender_result_for,
    verify,
)

TEAM = """Таблица специалистов
| Ф. И. О. | Должность |
|---|---|
| Иванова Анна Петровна | Партнер практики «Цифровая трансформация» |
| Петров Олег Ильич | Директор практики «Недвижимость» |
| Сидоров Иван Олегович | Руководитель проектов практики «Цифровая трансформация» |
| Козлов Пётр Ильич | Аналитик практики «Цифровая трансформация» |"""

CONTEXT = DraftContext(
    text="",
    slides={
        1: "Отчет от 18 июля 2025 г. Заказчик: ООО «Ромашка»",
        4: TEAM,
        5: "Договор от 31.03.2025. Выручка 129 млрд руб.",
    },
    titles={1: "Титул", 4: "Специалисты", 5: "Основание", 36: "Проведено 47 интервью"},
)


def answer(**fields: object) -> str:
    return json.dumps(
        {name: {"value": value, "slides": slides} for name, (value, slides) in fields.items()},
        ensure_ascii=False,
    )


def drafted(**fields: object):
    return verify(parse(answer(**fields)), CONTEXT)


def test_parse_rejects_non_json() -> None:
    with pytest.raises(ValueError, match="не JSON"):
        parse("Вот карточка проекта: ...")


def test_parse_ignores_unknown_and_fills_missing() -> None:
    fields = parse(answer(**{"Год": (2025, [1]), "Выдуманное поле": ("x", [1])}))
    assert "Выдуманное поле" not in fields
    assert fields["Год"].value == 2025
    assert fields["Клиент"].value is None


def test_supported_values_pass() -> None:
    fields = drafted(**{
        "Год": (2025, [1]),
        "Дата начала": ("2025-03-31", [5]),
        "Дата окончания": ("2025-07-18", [1]),  # «18 июля 2025» прописью
        "Клиент (по договору)": ("ООО «Ромашка»", [1]),
        "Описание клиента": ("Выручка 129 млрд руб.", [5]),
        "Тип документа": ("report", [1]),
    })
    assert all(f.verified for f in fields.values()), [f.problems for f in fields.values()]


def test_title_is_a_valid_source() -> None:
    # У отчётов заголовок слайда — вывод с цифрами; модель видит все заголовки.
    assert drafted(**{"Подход": ("47 интервью", [36])})["Подход"].verified


def test_slide_not_in_document_flagged() -> None:
    [problem] = drafted(**{"Подход": ("Интервью", [99])})["Подход"].problems[:1]
    assert "99" in problem


def test_name_typo_flagged() -> None:
    field = drafted(**{"Команда": ("Козлов Пётр Ильинич", [4])})["Команда"]
    assert any("Ильинич" in p for p in field.problems)


def test_wrong_date_number_and_enum_flagged() -> None:
    fields = drafted(**{
        "Дата начала": ("2025-04-01", [5]),
        "Описание клиента": ("Выручка 130 млрд руб.", [5]),
        "Тип документа": ("отчёт", [1]),
        "Год": (2024, [1]),
    })
    for name in ("Дата начала", "Описание клиента", "Тип документа", "Год"):
        assert not fields[name].verified, name


def test_generative_field_needs_no_source() -> None:
    field = drafted(**{"Обзелич.название клиента": ("крупный ритейлер", [])})
    assert field["Обзелич.название клиента"].verified


def test_roles_from_team_table_override_model() -> None:
    model = parse(answer(**{
        "Курирующий партнер": ("Петров Олег Ильич", [4]),  # модель перепутала роли
        "Команда": ("Иванова Анна Петровна", [4]),
    }))
    fields = apply_team_table(model, CONTEXT)
    assert fields["Курирующий партнер"].value == "Иванова Анна Петровна"
    assert fields["Директор"].value == "Петров Олег Ильич"
    assert fields["Менеджер проекта"].value == "Сидоров Иван Олегович"
    assert fields["Команда"].value == "Козлов Пётр Ильич"
    assert fields["Практика"].value == "Цифровая трансформация; Недвижимость"
    assert fields["Директор"].slides == [4]


def test_without_team_table_model_answer_kept() -> None:
    context = DraftContext(text="", slides={1: "без таблицы"}, titles={1: "Титул"})
    fields = apply_team_table(parse(answer(**{"Директор": ("Петров", [1])})), context)
    assert fields["Директор"].value == "Петров"


@pytest.mark.parametrize(
    ("doc_types", "expected"),
    [(["report"], "won"), (["kp", "act"], "won"), (["kp"], None), ([], None)],
)
def test_tender_result_rule(doc_types: list[str], expected: str | None) -> None:
    assert tender_result_for(doc_types) == expected


DECK = """---
doc_title: "Отчёт"
slides_total: 4
---
<!-- slide: 1 -->
## [Слайд 1] Титул
**Раздел:** Титул
Заказчик

<!-- slide: 2 -->
## [Слайд 2] Общие термины и сокращения
**Раздел:** Сведения
ЦТ — цифровая трансформация

<!-- slide: 3 -->
## [Слайд 3] Сведения о специалистах
**Раздел:** Сведения
""" + TEAM + """

<!-- slide: 4 -->
## [Слайд 4] Рынок вырос на 12%
**Раздел:** Анализ
Таблица рынка
"""


def test_context_takes_passport_slides_skips_glossary() -> None:
    deck = chunk_slides(DECK, lambda s: len(s.split()), hard_limit=8192)
    context = select_context(deck, lambda s: len(s.split()), budget=1000)
    assert 1 in context.slides and 3 in context.slides
    assert 2 not in context.slides  # глоссарий — пустая трата окна
    assert context.titles[4] == "Рынок вырос на 12%"
    assert "4. Рынок вырос на 12%" in context.text


def test_context_respects_budget() -> None:
    deck = chunk_slides(DECK, lambda s: len(s.split()), hard_limit=8192)
    context = select_context(deck, lambda s: len(s.split()), budget=25)
    assert sum(len(t.split()) for t in context.slides.values()) <= 25


openpyxl = pytest.importorskip("openpyxl")


def make_registry(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    projects = wb.active
    projects.title = "Проекты"
    projects.append(["Год", "Код проекта", "Клиент", "Название", "Дата начала", "Дата окончания",
                     "Продоожительность, нед", "Уровень публичности", "Результат тендера",
                     "Статус записи"])
    docs = wb.create_sheet("Документы")
    docs.append(["Код проекта", "Файл", "Исходный файл", "Тип документа", "Название",
                 "Дата документа", "Слайдов", "Комментарии"])
    wb.save(path)


def test_append_draft_row_and_document(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from bzpo.registry.workbook import DocumentRow, RegistryError, append_draft

    registry = tmp_path / "registry.xlsx"
    make_registry(registry)
    fields = drafted(**{
        "Год": (2025, [1]),
        "Клиент": ("ООО «Ромашка»", [1]),
        "Дата начала": ("2025-03-31", [5]),
        "Дата окончания": ("2025-04-01", [5]),  # не подтвердится -> оранжевая
    })
    document = DocumentRow("samples/x.md", None, "report", "Отчёт", date(2025, 7, 18), 4)
    row, code = append_draft(registry, fields, document, tender_result="won", model="giga")

    assert (row, code) == (2, "РОМАШКА-2025-DRAFT")
    ws = load_workbook(registry)["Проекты"]
    values = {ws.cell(1, c).value: ws.cell(2, c) for c in range(1, ws.max_column + 1)}
    assert values["Статус записи"].value == "черновик"
    assert values["Результат тендера"].value == "won"
    assert values["Уровень публичности"].value is None  # решает человек
    assert values["Дата начала"].value.date() == date(2025, 3, 31)
    assert values["Клиент"].fill.start_color.rgb.endswith("FFFF00")
    assert values["Дата окончания"].fill.start_color.rgb.endswith("F4B183")
    assert "НЕ ПОДТВЕРЖДЕНО" in values["Дата окончания"].comment.text
    assert values["Продоожительность, нед"].value.startswith("=IF(")
    doc = load_workbook(registry)["Документы"]
    assert [doc.cell(2, c).value for c in (1, 2, 4)] == [code, "samples/x.md", "report"]

    with pytest.raises(RegistryError, match="уже в реестре"):
        append_draft(registry, fields, document, tender_result="won", model="giga")


def test_nfd_filename_matches_registered_nfc(tmp_path: Path) -> None:
    import unicodedata

    from bzpo.registry.workbook import (
        DocumentRow,
        RegistryError,
        append_draft,
        ensure_not_registered,
        registered_files,
    )

    registry = tmp_path / "registry.xlsx"
    make_registry(registry)
    nfc = "samples/аналитический.md"
    nfd = unicodedata.normalize("NFD", nfc)
    assert nfc != nfd
    fields = drafted(**{"Год": (2025, [1])})
    append_draft(registry, fields, DocumentRow(nfd, None, "report", "Отчёт", None, 1),
                 tender_result="won", model="giga")
    assert nfc in registered_files(registry)  # хранится в NFC
    with pytest.raises(RegistryError, match="уже в реестре"):
        ensure_not_registered(registry, nfc)


def test_parse_tolerates_model_field_names_and_null_strings() -> None:
    raw = json.dumps({
        "год": {"value": "2025", "slides": [1]},  # регистр
        "обобщ.название клиента": {"value": "крупный ритейлер", "slides": []},  # опечатка
        "Директор": {"value": "null", "slides": [4]},  # «нет данных» строкой
        "КЛИЕНТ": {"value": "ООО «Ромашка»", "slides": [1]},
    }, ensure_ascii=False)
    fields = parse(raw)
    assert fields["Год"].value == "2025"
    assert fields["Обзелич.название клиента"].value == "крупный ритейлер"
    assert fields["Клиент"].value == "ООО «Ромашка»"
    assert fields["Директор"].value is None


def test_doc_type_from_text_overrides_model() -> None:
    from bzpo.registry.draft import guess_doc_type

    kp = ("Сопроводительное письмо. Мы понимаем, что заказчику нужно выбрать платформу. "
          "Мы предлагаем этапы работ. Стоимость работ — 3 900 000 руб. Срок выполнения — "
          "5 недель. По итогам — отчет.")
    assert guess_doc_type(kp) == "kp"
    assert guess_doc_type("Отчет. Результаты услуг. Сведения об исполнителе. Отчет.") == "report"
    assert guess_doc_type("Нейтральный текст без признаков") is None

    context = DraftContext(text="", slides={1: kp}, titles={1: "Предложение"})
    fields = verify(parse(answer(**{"Тип документа": ("report", [1])})), context)
    assert fields["Тип документа"].value == "kp"
    assert any("поставлено «kp»" in p for p in fields["Тип документа"].problems)
