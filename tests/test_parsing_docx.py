"""Автоопределение формата, Word -> md, чанкинг md-документа."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

docx_lib = pytest.importorskip("docx")

from bzpo.ingest.document import chunk_document  # noqa: E402
from bzpo.parsing.convert import convert, sources  # noqa: E402
from bzpo.parsing.detect import detect  # noqa: E402


def png(width: int, height: int) -> io.BytesIO:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def build_docx(path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("Отчёт о диагностике", style="Title")
    document.add_heading("1. Подход", level=1)
    document.add_paragraph("Проведено 47 интервью с руководством заказчика.")
    document.add_paragraph("Интервью топ-менеджмента", style="List Bullet")
    document.add_paragraph("Опрос 250 сотрудников", style="List Bullet")
    document.add_heading("1.1. Оценка", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Критерий", "Балл"
    table.cell(1, 0).text, table.cell(1, 1).text = "Данные", "2,1"
    document.add_picture(png(160, 90), width=Inches(6))  # широкий 16:9 — скриншот слайда
    document.add_paragraph("Рисунок 1. Сводная оценка зрелости", style="Caption")
    document.add_picture(png(100, 100), width=Inches(1))  # маленький квадрат — не слайд
    target = path / "report.docx"
    document.save(str(target))
    return target


def test_detect_by_content_not_extension(tmp_path: Path) -> None:
    docx_file = build_docx(tmp_path)
    disguised = tmp_path / "report.bin"
    disguised.write_bytes(docx_file.read_bytes())
    assert detect(disguised).format == "docx"
    (tmp_path / "a.pdf").write_bytes(b"%PDF-1.7 ...")
    assert detect(tmp_path / "a.pdf").format == "pdf"
    (tmp_path / "old.ppt").write_bytes(b"\xd0\xcf\x11\xe0" + b"0" * 10)
    assert detect(tmp_path / "old.ppt").format == "legacy_office"
    (tmp_path / "s.md").write_text("<!-- slide: 1 -->\n## [Слайд 1] Т", encoding="utf-8")
    assert detect(tmp_path / "s.md").format == "md_slides"
    (tmp_path / "d.md").write_text("# Документ\n\nТекст", encoding="utf-8")
    assert detect(tmp_path / "d.md").format == "md_document"
    (tmp_path / "~$report.docx").write_bytes(b"lock")
    assert all(not p.name.startswith("~$") for p in sources(tmp_path))


def test_docx_to_markdown(tmp_path: Path) -> None:
    result = convert(build_docx(tmp_path))
    assert result.detection.format == "docx" and result.ok
    md = result.markdown
    assert 'doc_title: "Отчёт о диагностике"' in md
    assert "## 1. Подход" in md and "### 1.1. Оценка" in md
    assert "- Интервью топ-менеджмента" in md
    assert "| Данные | 2,1 |" in md
    assert "<!-- figure: 1; kind: slide_screenshot -->" in md
    assert "<!-- figure: 2; kind: image -->" in md
    assert "*Подпись:* Рисунок 1. Сводная оценка зрелости" in md
    assert result.needs_vlm == 1  # маленький квадратный рисунок (логотип) не распознаётся
    assert "slide_screenshots: 1" in md


def test_unsupported_format_explained(tmp_path: Path) -> None:
    (tmp_path / "a.pdf").write_bytes(b"%PDF-1.7")
    result = convert(tmp_path / "a.pdf")
    assert not result.ok and "pdf" in (result.error or "")


def words(text: str) -> int:
    return len(text.split())


def test_document_chunks_follow_headings_and_keep_atomic_blocks(tmp_path: Path) -> None:
    md = convert(build_docx(tmp_path)).markdown
    deck = chunk_document(md, words, target=500, minimum=1, hard_limit=8192,
                          customer_label="крупный девелопер")
    assert deck.kind == "document"
    sections = [(c.section, c.subsection) for c in deck.chunks]
    assert ("1. Подход", None) in sections and ("1. Подход", "1.1. Оценка") in sections
    table = next(c for c in deck.chunks if c.is_table)
    assert table.text_raw.startswith("1.1. Оценка") and "| Данные | 2,1 |" in table.text_raw
    figure = next(c for c in deck.chunks if c.slide_type == "рисунок"
                  and "slide_screenshot" in c.text_raw)
    assert "скриншот слайда" in figure.text_enriched
    assert all(c.page is None for c in deck.chunks)
    assert "заказчик крупный девелопер" in deck.chunks[0].text_enriched
    assert "Отчёт о диагностике" in deck.chunks[0].text_enriched


def test_document_target_splits_long_section() -> None:
    paragraphs = "\n\n".join(f"Абзац {i} " + "слово " * 40 for i in range(6))
    md = f"---\ndoc_title: \"Д\"\n---\n# Д\n\n## Раздел\n\n{paragraphs}\n"
    deck = chunk_document(md, words, target=100, minimum=10, hard_limit=8192)
    assert len(deck.chunks) == 3  # по 2 абзаца (~84 слова) на чанк
    assert all(c.section == "Раздел" for c in deck.chunks)


class FakeOCR:
    model = "fake-ocr"

    def __init__(self, text: str) -> None:
        self.text = text
        self.calls = 0

    def generate(self, prompt: str, **kwargs: object) -> object:
        self.calls += 1
        assert kwargs["images"]

        class Result:
            text = self.text

        return Result()


def test_ocr_cleans_duplicates_and_noise(tmp_path: Path) -> None:
    from bzpo.parsing import ocr

    raw = ("Сводная оценка зрелости\n\nБалл 2,1 из 5\n\n```markdown\n"
           "Сводная оценка зрелости\n\nБалл 2,1 из 5\n```")
    client = FakeOCR(raw)
    assert ocr.transcribe(b"png", client, tmp_path, "m") == "Сводная оценка зрелости\nБалл 2,1 из 5"
    looping = "\n".join(["ОЦЕНКА ЗРЕЛОСТИ", "ПАО «Компания»"] * 8)  # зацикливание модели
    assert ocr.clean(looping) == "ОЦЕНКА ЗРЕЛОСТИ\nПАО «Компания»"
    assert ocr.transcribe(b"png", client, tmp_path, "m") and client.calls == 1  # из кэша
    assert ocr.transcribe(b"logo", FakeOCR("COMPANY"), tmp_path, "m") == ""  # шум


def test_docx_ocr_text_replaces_placeholder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from bzpo.parsing import convert as convert_module

    class Client(FakeOCR):
        def __enter__(self) -> FakeOCR:
            return self

        def __exit__(self, *exc: object) -> None: ...

    monkeypatch.setattr("bzpo.providers.ollama.OllamaClient",
                        lambda config, role: Client("Сводная оценка\n\nБалл 2,1 из 5"))
    result = convert_module.convert(build_docx(tmp_path), vlm=True, cache_dir=tmp_path / "c")
    md = result.markdown
    assert "**Рисунок 1 (скриншот слайда; текст распознан OCR — проверить):**" in md
    assert "Сводная оценка\nБалл 2,1 из 5" in md
    assert result.needs_vlm == 0 and "текст распознан на 1" in result.notes
    deck = chunk_document(md, words, target=500, minimum=1, hard_limit=8192)
    figure = next(c for c in deck.chunks if "Сводная оценка" in c.text_raw)
    assert "Балл 2,1 из 5" in figure.text_raw  # рисунок и его текст — один чанк


def test_one_failing_figure_does_not_stop_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from bzpo.parsing import convert as convert_module
    from bzpo.providers.ollama import OllamaUnavailable

    class Failing(FakeOCR):
        def generate(self, prompt: str, **kwargs: object) -> object:
            raise OllamaUnavailable("ReadTimeout: timed out")

        def __enter__(self) -> FakeOCR:
            return self

        def __exit__(self, *exc: object) -> None: ...

    monkeypatch.setattr("bzpo.providers.ollama.OllamaClient",
                        lambda config, role: Failing(""))
    result = convert_module.convert(build_docx(tmp_path), vlm=True, cache_dir=tmp_path / "c")
    assert result.markdown is not None  # документ сконвертирован
    assert any("модель не ответила" in issue for issue in result.vlm_issues)
    assert result.needs_vlm == 1  # рисунок остался на разбор
