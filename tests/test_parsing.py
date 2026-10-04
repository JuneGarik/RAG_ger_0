"""Конвертер pptx -> md и проверка md против XML — на презентации, собранной в тесте."""

from __future__ import annotations

from pathlib import Path

import pytest

pptx_lib = pytest.importorskip("pptx")

from bzpo.ingest.slides import chunk_slides  # noqa: E402
from bzpo.parsing.check import check  # noqa: E402
from bzpo.parsing.pptx import UNNAMED_CHART_NOTE, extract, to_markdown  # noqa: E402


def build(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    title_only, title_body = prs.slide_layouts[5], prs.slide_layouts[1]

    divider = prs.slides.add_slide(title_only)  # 1: разделитель «01»
    divider.shapes.title.text = "Стратегические цели"
    divider.shapes.add_textbox(Inches(1), Inches(4), Inches(1), Inches(1)).text_frame.text = "01"

    text = prs.slides.add_slide(title_body)  # 2: текст со списком и ссылкой
    text.shapes.title.text = "Выручка выросла на 12%"
    frame = text.placeholders[1].text_frame
    frame.text = "Проведено 47 интервью с руководством"
    sub = frame.add_paragraph()
    sub.text = "Опрос охватил 250 сотрудников"
    sub.level = 1
    run = frame.add_paragraph().add_run()
    run.text = "Отчёт АКРА"
    run.hyperlink.address = "https://example.org/report/1689928602786"

    table_slide = prs.slides.add_slide(title_only)  # 3: таблица с ячейкой-цветом
    table_slide.shapes.title.text = "Тепловая карта зрелости"
    table = table_slide.shapes.add_table(2, 3, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for i, value in enumerate(["Критерий", "Продажи", "ИТ"]):
        table.cell(0, i).text = value
    table.cell(1, 0).text = "Данные"
    table.cell(1, 1).fill.solid()
    table.cell(1, 1).fill.fore_color.rgb = RGBColor(0x54, 0xCF, 0xD6)
    table.cell(1, 2).text = "2,1"

    chart_slide = prs.slides.add_slide(title_only)  # 4: диаграмма
    chart_slide.shapes.title.text = "Финансовые показатели"
    data = CategoryChartData()
    data.categories = ["2023", "2024"]
    data.add_series("Выручка", (146764.95, 128649.79))
    chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2),
                                 Inches(6), Inches(4), data)
    target = path / "deck.pptx"
    prs.save(str(target))
    return target


def test_extract_structure_and_markdown(tmp_path: Path) -> None:
    deck = extract(build(tmp_path))
    assert [s.number for s in deck.slides] == [1, 2, 3, 4]
    assert deck.slides[0].section == "01. Стратегические цели"
    assert deck.slides[0].slide_type == "разделитель раздела"
    assert deck.slides[3].section == "01. Стратегические цели"  # раздел тянется дальше

    md = to_markdown(deck)
    assert "Проведено 47 интервью с руководством" in md
    assert "  - Опрос охватил 250 сотрудников" not in md  # уровень 1 -> без отступа
    assert "- Опрос охватил 250 сотрудников" in md
    assert "[Отчёт АКРА](https://example.org/report/1689928602786)" in md
    assert "| Данные | ■ #54CFD6 | 2,1 |" in md  # значение цветом сохранено
    assert deck.slides[2].needs_vlm and not deck.slides[1].needs_vlm
    assert "| 2023 | 146764,95 |" in md and "| Категория | Выручка |" in md
    assert UNNAMED_CHART_NOTE not in md  # у этой диаграммы подписи есть в XML
    assert "converter_version: pptx-xml-1" in md


def test_output_readable_by_indexer(tmp_path: Path) -> None:
    md = to_markdown(extract(build(tmp_path)))
    chunks = chunk_slides(md, lambda s: len(s.split()), hard_limit=8192).chunks
    assert [c.slide for c in chunks] == [1, 2, 3, 4]
    assert chunks[1].section == "01. Стратегические цели"
    assert "Источник значений" not in chunks[1].text_raw


def test_check_passes_own_output_and_catches_defects(tmp_path: Path) -> None:
    pptx = build(tmp_path)
    md = to_markdown(extract(pptx))
    result = check(pptx, md)
    assert result.passed, (result.coverage, [(s.slide, s.unsupported) for s in result.slides])

    invented = md.replace("Проведено 47 интервью", "Проведено 52 интервью")
    assert check(pptx, invented).slides[1].unsupported == ["52"]

    rounded = md.replace("146764,95", "146 765")  # округление значения диаграммы — допустимо
    assert check(pptx, rounded).unsupported == 0

    dropped = md.replace("Проведено 47 интервью с руководством", "")
    assert check(pptx, dropped).slides[1].coverage < 1.0

    one_less = md.split("<!-- slide: 4 -->")[0]
    assert "слайдов в md 3, в pptx 4" in check(pptx, one_less).problems
