"""Запись черновика в реестр Excel (data/registry.xlsx).

Черновик добавляется строкой в лист «Проекты» со статусом «черновик» и
строкой в лист «Документы». Каждая заполненная ячейка — жёлтая с
примечанием, откуда взято значение; поле, не прошедшее проверку, —
оранжевое с причиной. Загрузчик реестра берёт только строки со статусом
«проверено»: статус меняет человек после просмотра.

openpyxl импортируется лениво: он нужен только этой операции
(`pip install -e ".[registry]"`).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from bzpo.registry.draft import DraftField

PROJECTS = "Проекты"
DOCUMENTS = "Документы"
STATUS_COLUMN = "Статус записи"
STATUS_DRAFT = "черновик"
STATUS_VERIFIED = "проверено"
_YELLOW = "FFFF00"
_ORANGE = "F4B183"
_AUTHOR = "БЗПО"
# Не годятся в код проекта: «ООО «ДЕВ «Лидер»» -> ДЕВ, а не ООО.
_LEGAL_FORMS = {"ООО", "АО", "ПАО", "ЗАО", "ОАО", "НАО", "ГК", "ИП", "ФГУП", "ГУП", "МУП",
                "LLC", "JSC", "LTD", "INC"}


def file_key(path: str) -> str:
    """Путь файла для сравнения и хранения: NFC.

    macOS отдаёт имена файлов в NFD («й» = «и» + знак краткости), а набранный
    текст и база — в NFC. Без нормализации один и тот же файл не находится в
    реестре и вносится повторно.
    """
    return unicodedata.normalize("NFC", path)


class RegistryError(ValueError):
    """Реестр не той структуры или документ в нём уже есть."""


@dataclass(frozen=True, slots=True)
class DocumentRow:
    file: str  # путь относительно data/
    source_file: str | None
    doc_type: str | None
    title: str | None
    doc_date: date | None
    slides: int | None


def append_draft(
    path: Path,
    fields: dict[str, DraftField],
    document: DocumentRow,
    *,
    tender_result: str | None,
    model: str,
) -> tuple[int, str]:
    """Добавить черновик; вернуть номер строки на листе «Проекты» и временный код."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = load_workbook(path)
    projects = _sheet(workbook, PROJECTS)
    documents = _sheet(workbook, DOCUMENTS)
    project_cols = _columns(projects)
    document_cols = _columns(documents)
    for required in ("Код проекта", STATUS_COLUMN, "Результат тендера"):
        if required not in project_cols:
            raise RegistryError(f"в листе «{PROJECTS}» нет колонки «{required}»")

    _ensure_new(documents, document_cols, document.file)

    code = _draft_code(fields, document)
    row = _first_empty_row(projects, project_cols["Код проекта"], project_cols["Название"])
    font = Font(name="Arial")
    wrap = Alignment(wrap_text=True, vertical="top")

    def put(column: int, value: Any, note: str | None, fill: str | None) -> None:
        cell = projects.cell(row, column, value)
        cell.font = font
        cell.alignment = wrap
        if isinstance(value, date):
            cell.number_format = "DD.MM.YYYY"
        if fill:
            cell.fill = PatternFill("solid", start_color=fill)
        if note:
            cell.comment = Comment(note, _AUTHOR)

    for name, draft in fields.items():
        if name not in project_cols or draft.value is None:
            continue
        value: Any = draft.value
        if name.startswith("Дата"):
            value = _as_date(value) or value
        if name == "Год" and str(value).isdigit():
            value = int(value)
        slides = ", ".join(map(str, draft.slides)) or "—"
        if draft.verified:
            put(project_cols[name], value, f"Черновик ({model}). Слайды: {slides}", _YELLOW)
        else:
            note = f"НЕ ПОДТВЕРЖДЕНО ({model}): {'; '.join(draft.problems)}. Слайды: {slides}"
            put(project_cols[name], value, note, _ORANGE)

    put(project_cols["Код проекта"], code,
        "Временный код. Меняя его, поменяйте и на листе «Документы»", _YELLOW)
    put(project_cols[STATUS_COLUMN], STATUS_DRAFT,
        f"Проверьте строку и поставьте «{STATUS_VERIFIED}»: загрузчик берёт только их", _YELLOW)
    if tender_result:
        put(project_cols["Результат тендера"], tender_result,
            "Правило: в проекте есть отчёт или акт — тендер выигран", None)
    else:
        put(project_cols["Результат тендера"], None,
            "Только КП: укажите won / lost / not_submitted", _YELLOW)
    if "Уровень публичности" in project_cols:
        put(project_cols["Уровень публичности"], None,
            "Заполняет человек: open / anon / closed", _YELLOW)
    duration = project_cols.get("Продоожительность, нед")
    start, end = project_cols.get("Дата начала"), project_cols.get("Дата окончания")
    if duration and start and end:
        s, e = _letter(start), _letter(end)
        put(duration, f'=IF(AND(ISNUMBER({s}{row}),ISNUMBER({e}{row})),'
                      f'ROUND(({e}{row}-{s}{row})/7,0),"")', None, None)

    _write_document(documents, document_cols, code, document,
                    f"Добавлен черновиком {date.today():%d.%m.%Y}")
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(path)
    return row, code


def add_document(path: Path, code: str, document: DocumentRow) -> None:
    """Добавить документ к существующему проекту (строка листа «Документы»).

    Проект уже проверен человеком, поэтому строка проекта не трогается: новый
    документ (КП, акт, ещё один отчёт) попадает в индекс при следующем прогоне.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    projects = _sheet(workbook, PROJECTS)
    documents = _sheet(workbook, DOCUMENTS)
    project_cols, document_cols = _columns(projects), _columns(documents)
    codes = {projects.cell(r, project_cols["Код проекта"]).value
             for r in range(2, projects.max_row + 1)}
    if code not in codes:
        raise RegistryError(f"проекта {code} нет в реестре")
    _ensure_new(documents, document_cols, document.file)
    _write_document(documents, document_cols, code, document,
                    f"Добавлен из интерфейса {date.today():%d.%m.%Y}")
    workbook.save(path)


def _write_document(documents: Any, columns: dict[str, int], code: str,
                    document: DocumentRow, comment: str) -> None:
    from openpyxl.styles import Alignment, Font

    row = documents.max_row + 1
    values = {
        "Код проекта": code, "Файл": file_key(document.file),
        "Исходный файл": file_key(document.source_file) if document.source_file else None,
        "Тип документа": document.doc_type, "Название": document.title,
        "Дата документа": document.doc_date, "Слайдов": document.slides,
        "Комментарии": comment,
    }
    for name, value in values.items():
        if name in columns:
            cell = documents.cell(row, columns[name], value)
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"


def registered_files(path: Path) -> dict[str, tuple[str | None, str | None]]:
    """Файл (относительно data/) -> (код проекта, статус строки проекта)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    projects = _sheet(workbook, PROJECTS)
    documents = _sheet(workbook, DOCUMENTS)
    pcols, dcols = _columns(projects), _columns(documents)
    status = {
        row[pcols["Код проекта"] - 1]: row[pcols[STATUS_COLUMN] - 1]
        for row in projects.iter_rows(min_row=2, values_only=True)
        if STATUS_COLUMN in pcols
    }
    result = {}
    for row in documents.iter_rows(min_row=2, values_only=True):
        file, code = row[dcols["Файл"] - 1], row[dcols["Код проекта"] - 1]
        if file:
            result[file_key(str(file))] = (code, status.get(code))
    return result


def ensure_not_registered(path: Path, file: str) -> None:
    """Проверить до генерации: документ, уже внесённый в реестр, не обрабатывается."""
    from openpyxl import load_workbook

    documents = _sheet(load_workbook(path, read_only=True), DOCUMENTS)
    _ensure_new(documents, _columns(documents), file)


# --------------------------------------------------------------- детали


def _ensure_new(documents: Any, columns: dict[str, int], file: str) -> None:
    for row in documents.iter_rows(min_row=2, values_only=True):
        if row[columns["Файл"] - 1] and file_key(str(row[columns["Файл"] - 1])) == file_key(file):
            raise RegistryError(f"{file} уже в реестре (проект {row[columns['Код проекта'] - 1]})")



def _sheet(workbook: Any, name: str) -> Any:
    if name not in workbook.sheetnames:
        raise RegistryError(f"в реестре нет листа «{name}»")
    return workbook[name]


def _columns(sheet: Any) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None
    }


def _first_empty_row(sheet: Any, *key_columns: int) -> int:
    row = 2
    while any(sheet.cell(row, col).value not in (None, "") for col in key_columns):
        row += 1
    return row


def _draft_code(fields: dict[str, DraftField], document: DocumentRow) -> str:
    client = fields.get("Клиент")
    year = fields.get("Год")
    base = str(client.value) if client and client.value else Path(document.file).stem
    words = [
        w for w in re.findall(r"[A-ZА-ЯЁ]{2,}|[A-Za-zА-Яа-яЁё]{3,}", base)
        if w.upper() not in _LEGAL_FORMS
    ]
    slug = (words[0] if words else "ПРОЕКТ").upper()[:12]
    suffix = str(year.value) if year and year.value else f"{date.today():%Y}"
    return f"{slug}-{suffix}-DRAFT"


def _as_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def _letter(column: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(column)


def is_registry(path: Path) -> bool:
    """xlsx — реестр БЗПО: есть листы «Проекты» и «Документы»."""
    from openpyxl import load_workbook

    try:
        names = load_workbook(path, read_only=True).sheetnames
    except Exception:  # noqa: BLE001 - любой нечитаемый файл — не реестр
        return False
    return PROJECTS in names and DOCUMENTS in names


def review_project(
    path: Path,
    code: str,
    *,
    new_code: str,
    publicity: str,
    anon_label: str | None,
    doc_type: str | None,
    tender_result: str | None,
    verified: bool,
) -> None:
    """Записать решения человека по строке проекта (форма проверки в интерфейсе).

    Код меняется на обоих листах сразу — иначе документы теряют проект. Жёлтая
    заливка снимается с исправленных ячеек: их проверил человек.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    workbook = load_workbook(path)
    projects, documents = _sheet(workbook, PROJECTS), _sheet(workbook, DOCUMENTS)
    pcols, dcols = _columns(projects), _columns(documents)
    row = next((r for r in range(2, projects.max_row + 1)
                if projects.cell(r, pcols["Код проекта"]).value == code), None)
    if row is None:
        raise RegistryError(f"проекта {code} нет в реестре")
    if new_code != code and any(projects.cell(r, pcols["Код проекта"]).value == new_code
                                for r in range(2, projects.max_row + 1)):
        raise RegistryError(f"код {new_code} уже занят другим проектом")
    no_fill = PatternFill(fill_type=None)
    values = {
        "Код проекта": new_code, "Уровень публичности": publicity,
        "Обзелич.название клиента": anon_label, "Результат тендера": tender_result,
        STATUS_COLUMN: STATUS_VERIFIED if verified else STATUS_DRAFT,
    }
    for name, value in values.items():
        if name in pcols and (value is not None or name != "Обзелич.название клиента"):
            cell = projects.cell(row, pcols[name], value)
            cell.fill = no_fill
    for r in range(2, documents.max_row + 1):
        if documents.cell(r, dcols["Код проекта"]).value == code:
            documents.cell(r, dcols["Код проекта"], new_code)
            if doc_type and "Тип документа" in dcols:
                documents.cell(r, dcols["Тип документа"], doc_type)
    workbook.save(path)


def project_row(path: Path, code: str) -> dict[str, Any]:
    """Значения строки проекта по именам колонок (для формы проверки)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    projects, documents = _sheet(workbook, PROJECTS), _sheet(workbook, DOCUMENTS)
    header = [str(c.value).strip() if c.value is not None else "" for c in projects[1]]
    for values in projects.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values, strict=False))
        if row.get("Код проекта") == code:
            dheader = [str(c.value).strip() if c.value is not None else "" for c in documents[1]]
            row["_documents"] = [dict(zip(dheader, v, strict=False))
                                 for v in documents.iter_rows(min_row=2, values_only=True)
                                 if v and v[dheader.index("Код проекта")] == code]
            return row
    raise RegistryError(f"проекта {code} нет в реестре")
