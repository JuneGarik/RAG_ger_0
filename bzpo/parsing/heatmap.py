"""Значения тепловых карт по цвету заливки и легенде слайда — кодом, без модели.

Тепловая карта в отчётах — таблица, где значение ячейки передано цветом, а
расшифровка дана легендой: группа «цветной прямоугольник + подпись»
(«≤ 1.7», «1.8-1.9», …, «N/A»). Мультимодальная модель на такой карте
ошибается (на слайде 34 ДЕВ — 2 верных ячейки из 12), а XML даёт и цвета
ячеек, и цвета легенды точно.

Цвета ячеек и легенды не совпадают побайтно (ячейка #4ED1D8 при легенде
#02C0CA), поэтому сопоставление идёт по порядку яркости: оттенки ячеек
группируются, группы и позиции легенды упорядочиваются от тёмного к светлому.
Если число групп не равно числу позиций легенды — значения не подставляются
(слайд остаётся на разбор моделью), лучше честный пробел, чем сдвиг шкалы.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

THEME = "THEME"  # заливка цветом темы без RGB (так обычно рисуют «нет данных»)
_SAME_SHADE = 18  # расстояние в RGB, внутри которого оттенки — один цвет шкалы
_LABEL_MAX = 15


@dataclass(frozen=True, slots=True)
class LegendItem:
    label: str
    color: str  # RGB «RRGGBB» или THEME


def legend(shapes: Any, text_of: Any) -> list[LegendItem]:
    """Позиции легенды: группы ровно из одной залитой фигуры без текста и одной подписи."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    items: list[LegendItem] = []
    for shape in shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
            continue
        children = list(shape.shapes)
        texts = [c for c in children if text_of(c)]
        fills = [c for c in children if not text_of(c) and shape_fill(c) is not None]
        if len(children) == 2 and len(texts) == 1 and len(fills) == 1:
            label = text_of(texts[0])
            if len(label) <= _LABEL_MAX:
                items.append(LegendItem(label, shape_fill(fills[0]) or THEME))
        else:
            items += legend(shape.shapes, text_of)
    return items


def shape_fill(shape: Any) -> str | None:
    """RGB сплошной заливки, THEME для цвета темы, None — без сплошной заливки."""
    try:
        fill = shape.fill
        if fill.type != 1:  # MSO_FILL.SOLID
            return None
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        return None
    try:
        return str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return THEME


def map_cells(rows: list[list[str]], items: list[LegendItem]) -> list[list[str]] | None:
    """Заменить «■ #RRGGBB» / «■ THEME» в ячейках на подписи легенды; None — не удалось."""
    colored = [i for i in items if i.color != THEME]
    theme_items = [i for i in items if i.color == THEME]
    cell_colors = sorted({c[3:] for row in rows for c in row
                          if c.startswith("■ #")}, key=_luminance)
    if not colored or not cell_colors:
        return None
    clusters: list[list[str]] = []
    for color in cell_colors:
        if clusters and _distance(clusters[-1][0], color) <= _SAME_SHADE:
            clusters[-1].append(color)
        else:
            clusters.append([color])
    if len(clusters) != len(colored):
        return None
    ordered = sorted(colored, key=lambda i: _luminance(i.color))
    label_of = {color: item.label for cluster, item in zip(clusters, ordered, strict=True)
                for color in cluster}
    theme_label = theme_items[0].label if len(theme_items) == 1 else None

    def convert(cell: str, in_scale_row: bool) -> str:
        if cell.startswith("■ #"):
            return label_of[cell[3:]]
        if cell == f"■ {THEME}":
            # Цвет темы — «нет данных» только среди ячеек значений; в шапке
            # таблицы это просто оформление.
            return (theme_label or "") if in_scale_row else ""
        return cell

    return [[convert(c, any(x.startswith("■ #") for x in row)) for c in row] for row in rows]


def _rgb(color: str) -> tuple[int, int, int]:
    return int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)


def _luminance(color: str) -> float:
    r, g, b = _rgb(color)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _distance(a: str, b: str) -> float:
    return sum((x - y) ** 2 for x, y in zip(_rgb(a), _rgb(b), strict=True)) ** 0.5
