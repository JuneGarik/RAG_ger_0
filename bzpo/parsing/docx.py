"""Детерминированное извлечение docx -> md (docs/PARSING_PLAN.md, этап П3, без модели).

Из XML документа в исходном порядке: заголовки (по стилям Word) -> заголовки
md, абзацы, списки с уровнями, таблицы, подписи рисунков. Рисунки модель
не видит, поэтому каждый становится блоком

    <!-- figure: N; kind: slide_screenshot -->
    *[Рисунок N: подпись — содержимое не распознано; похоже на скриншот слайда]*

Скриншот слайда определяется по геометрии: широкий (≥ 60% ширины текста) и с
пропорциями экрана (от 4:3 до 16:9). Такой рисунок — полноценный слайд: его
транскрипция — задача прохода мультимодальной модели (П2), до неё документ
помечен `needs_vlm`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from bzpo.parsing.common import cell_text, clean, rows_to_markdown, yaml_string

CONVERTER_VERSION = "docx-xml-1"

_HEADING = re.compile(r"^(?:heading|заголовок)\s*(\d)$", re.I)
_TITLE_STYLES = {"title", "название", "заголовок"}
_CAPTION_STYLES = {"caption", "название объекта", "подпись"}
_SCREEN_RATIO = (1.25, 1.9)  # 4:3 = 1.33, 16:9 = 1.78
_WIDE_SHARE = 0.6


@dataclass(slots=True)
class Figure:
    number: int
    alt: str
    width: int
    height: int
    slide_screenshot: bool


@dataclass(slots=True)
class DocumentData:
    source: Path
    title: str
    blocks: list[str] = field(default_factory=list)
    figures: list[Figure] = field(default_factory=list)
    headings: int = 0

    @property
    def needs_vlm(self) -> int:
        return len(self.figures)


def extract(path: Path) -> DocumentData:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(path))
    section = document.sections[0] if document.sections else None
    text_width = (
        (section.page_width - section.left_margin - section.right_margin)
        if section is not None and section.page_width else None
    )
    data = DocumentData(path, "")
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._tbl: t for t in document.tables}
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p") and element in paragraphs:
            _paragraph(paragraphs[element], data, text_width)
        elif element.tag == qn("w:tbl") and element in tables:
            data.blocks.append(_table(tables[element]))

    core = (document.core_properties.title or "").strip()
    if not data.title:
        data.title = core or next(
            (b.lstrip("# ") for b in data.blocks if b.startswith("#")), path.stem
        )
    return data


def to_markdown(doc: DocumentData, *, source_name: str | None = None) -> str:
    screenshots = sum(f.slide_screenshot for f in doc.figures)
    front = [
        "---",
        f'doc_title: "{yaml_string(doc.title)}"',
        'doc_type: "Документ"',
        f'source_file: "{yaml_string(source_name or doc.source.name)}"',
        "source_format: docx",
        f"converter_version: {CONVERTER_VERSION}",
        f"converted_at: {date.today().isoformat()}",
        f"figures_total: {len(doc.figures)}",
        f"slide_screenshots: {screenshots}",
        f"needs_vlm: {doc.needs_vlm}",
        "---",
        "",
        f"# {doc.title}",
        "",
    ]
    body = [block for b in doc.blocks if b for block in (b, "")]
    return "\n".join(front + body)


# --------------------------------------------------------------- элементы


def _paragraph(paragraph: Any, data: DocumentData, text_width: int | None) -> None:
    style = (paragraph.style.name if paragraph.style is not None else "").strip()
    for drawing in paragraph._p.xpath(".//w:drawing"):
        _figure(drawing, data, text_width)
    text = clean(paragraph.text)
    if not text:
        return
    heading = _HEADING.match(style)
    if style.lower() in _TITLE_STYLES and not data.title:
        data.title = text
        return
    if heading:
        data.headings += 1
        data.blocks.append("#" * (min(int(heading.group(1)), 4) + 1) + " " + text)
        return
    if style.lower() in _CAPTION_STYLES:
        data.blocks.append(f"*Подпись:* {text}")
        return
    level = _list_level(paragraph)
    if level is not None:
        data.blocks.append("  " * level + "- " + text)
        return
    data.blocks.append(text)


def _list_level(paragraph: Any) -> int | None:
    levels = paragraph._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
    if levels:
        return int(levels[0])
    if paragraph._p.xpath("./w:pPr/w:numPr"):
        return 0
    style = (paragraph.style.name if paragraph.style is not None else "").lower()
    return 0 if "list" in style or "списо" in style else None


def _figure(drawing: Any, data: DocumentData, text_width: int | None) -> None:
    extent = drawing.xpath(".//wp:extent")
    width = int(extent[0].get("cx", 0)) if extent else 0
    height = int(extent[0].get("cy", 0)) if extent else 0
    alt = (drawing.xpath(".//wp:docPr/@descr") or drawing.xpath(".//wp:docPr/@name") or [""])[0]
    ratio = width / height if height else 0.0
    wide = bool(text_width) and width >= _WIDE_SHARE * text_width
    screenshot = wide and _SCREEN_RATIO[0] <= ratio <= _SCREEN_RATIO[1]
    number = len(data.figures) + 1
    data.figures.append(Figure(number, clean(alt), width, height, screenshot))
    kind = "slide_screenshot" if screenshot else "image"
    hint = "; похоже на скриншот слайда" if screenshot else ""
    data.blocks.append(
        f"<!-- figure: {number}; kind: {kind} -->\n"
        f"*[Рисунок {number}: {clean(alt) or 'без подписи'} — содержимое не распознано{hint}]*"
    )


def _table(table: Any) -> str:
    rows = []
    for row in table.rows:
        cells = []
        previous = None
        for cell in row.cells:
            # Объединённая ячейка в python-docx повторяется — выводим один раз.
            if previous is not None and cell._tc is previous:
                cells.append("")
                continue
            previous = cell._tc
            cells.append(cell_text(cell.text))
        rows.append(cells)
    return rows_to_markdown(rows)

